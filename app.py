import os
import re
import sqlite3
import shutil
from datetime import datetime
from collections import Counter
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, scrolledtext

# 外部ライブラリ
import openpyxl
import docx
from pypdf import PdfReader

DB_FILE = "light_index_v2.db"

SIGNAL_WORDS = ["回答", "結論", "判断", "結果", "対応", "不可", "可能", "不要", "必要", "原則", "特例", "承認", "却下", "につき", "ため、", "とおり"]
STOP_WORDS = {"こと", "もの", "ため", "これ", "それ", "あり", "なし", "する", "なる", "いる", "ある", "れる", "られる", "について", "における", "および", "または", "場合", "確認", "資料", "照会", "シート", "ページ"}

class UltraLightSearchAppV2:
    def __init__(self, root):
        self.root = root
        # 👇 タイトルを「交通局 資料検索＆提案システム」に変更
        self.root.title("交通局 資料検索＆提案システム")
        self.root.geometry("1000x780")

        # --- 1. 上部：フォルダ選択＆インデックス更新 ---
        frame_top = tk.Frame(root, pady=5)
        frame_top.pack(fill=tk.X, padx=10)
        
        btn_update = tk.Button(frame_top, text="① 資料フォルダを選択して記憶（初回・更新）", command=self.build_index, bg="#e1f5fe", font=("Meiryo", 9, "bold"))
        btn_update.pack(side=tk.LEFT, padx=5)
        
        self.lbl_status = tk.Label(frame_top, text="状態: 準備完了", fg="green", font=("Meiryo", 9))
        self.lbl_status.pack(side=tk.LEFT, padx=10)

        # 👇 フォルダ選択ボタンのすぐ下に注意書きを追加
        lbl_note = tk.Label(root, text="※インデックス作成中に画面が「（応答なし）」になることがありますが、裏で正常に処理中ですのでそのままお待ちください。", fg="#d32f2f", font=("Meiryo", 8, "italic"))
        lbl_note.pack(anchor="w", padx=15, pady=1)

        # --- 2. 検索バー ＆ ソート ---
        frame_search = tk.LabelFrame(root, text="② キーワード検索 ＆ 並び替え", pady=5, padx=5)
        frame_search.pack(fill=tk.X, padx=10, pady=5)

        self.ent_keyword = tk.Entry(frame_search, font=("Meiryo", 11))
        self.ent_keyword.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        self.ent_keyword.bind("<Return>", lambda event: self.run_search())

        # 👇 日付順ソート機能
        tk.Label(frame_search, text="並び順:", font=("Meiryo", 9)).pack(side=tk.LEFT, padx=5)
        self.sort_var = tk.StringVar(value="新着順")
        sort_combo = ttk.Combobox(frame_search, textvariable=self.sort_var, values=["新着順（新しい順）", "古い順", "関連度順"], width=15, state="readonly")
        sort_combo.pack(side=tk.LEFT, padx=5)
        sort_combo.bind("<<ComboboxSelected>>", lambda e: self.run_search())

        btn_search = tk.Button(frame_search, text="検索", command=self.run_search, bg="#c8e6c9", font=("Meiryo", 10, "bold"), width=10)
        btn_search.pack(side=tk.LEFT, padx=5)

        # --- 3. 提案エリア（関連ワード ＆ 要点ピックアップ） ---
        frame_suggest = tk.LabelFrame(root, text="💡 システムからの提案（要点まとめ ＆ 関連キーワード）", pady=5, padx=5, bg="#fffde7")
        frame_suggest.pack(fill=tk.X, padx=10, pady=5)

        self.frame_chips = tk.Frame(frame_suggest, bg="#fffde7")
        self.frame_chips.pack(fill=tk.X, padx=5, pady=2)

        self.txt_summary = scrolledtext.ScrolledText(frame_suggest, height=5, bg="#fff9c4", font=("Meiryo", 9))
        self.txt_summary.pack(fill=tk.X, padx=5, pady=3)

        # --- 4. 検索結果リスト（ファイル単位の折りたたみカード表示） ---
        tk.Label(root, text="📄 該当ファイル一覧（＋ボタンで簡易プレビュー・保存・元ファイルを開く）", font=("Meiryo", 9, "bold")).pack(anchor="w", padx=10)
        
        self.canvas = tk.Canvas(root, bg="#f5f5f5")
        self.scrollbar = ttk.Scrollbar(root, orient="vertical", command=self.canvas.yview)
        self.scrollable_frame = tk.Frame(self.canvas, bg="#f5f5f5")

        self.scrollable_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=10, pady=5)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def extract_text(self, filepath):
        ext = os.path.splitext(filepath)[1].lower()
        chunks = []
        try:
            if ext == ".xlsx":
                wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
                for sheet in wb.sheetnames:
                    s = wb[sheet]
                    sheet_lines = []
                    for row in s.iter_rows(values_only=True):
                        row_txt = " | ".join([str(cell).strip() for cell in row if cell is not None and str(cell).strip() != ""])
                        if row_txt:
                            sheet_lines.append(row_txt)
                    if sheet_lines:
                        full_sheet_txt = "\n".join(sheet_lines)
                        for i in range(0, len(full_sheet_txt), 700):
                            chunks.append((f"シート:{sheet}", full_sheet_txt[i:i+800]))
                wb.close()
            elif ext == ".docx":
                doc = docx.Document(filepath)
                full_txt = "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
                for i in range(0, len(full_txt), 700):
                    chunks.append(("Word本文", full_txt[i:i+800]))
            elif ext == ".pdf":
                with open(filepath, "rb") as f:
                    reader = PdfReader(f)
                    for idx, page in enumerate(reader.pages):
                        t = page.extract_text()
                        if t and t.strip():
                            chunks.append((f"P.{idx+1}", t.strip()))
            elif ext in [".txt", ".csv"]:
                with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                    full_txt = f.read()
                    for i in range(0, len(full_txt), 700):
                        chunks.append(("テキスト", full_txt[i:i+800]))
        except Exception:
            pass
        return chunks

    def build_index(self):
        folder = filedialog.askdirectory(title="照会関係資料フォルダを選択")
        if not folder:
            return
        self.lbl_status.config(text="状態: 高速インデックス作成中...", fg="blue")
        self.root.update()

        conn = sqlite3.connect(DB_FILE)
        c = conn.cursor()
        c.execute("DROP TABLE IF EXISTS docs")
        c.execute("CREATE VIRTUAL TABLE docs USING fts5(filename, location, filepath, content, mtime UNINDEXED, tokenize='trigram')")

        file_count = 0
        for root_dir, _, files in os.walk(folder):
            for f in files:
                if f.startswith("~$"):
                    continue
                if f.lower().endswith((".xlsx", ".docx", ".pdf", ".txt", ".csv")):
                    fullpath = os.path.join(root_dir, f)
                    try:
                        mtime = os.path.getmtime(fullpath)
                    except:
                        mtime = 0.0

                    chunks = self.extract_text(fullpath)
                    if chunks:
                        for loc, content in chunks:
                            c.execute("INSERT INTO docs VALUES (?, ?, ?, ?, ?)", (f, loc, fullpath, content, mtime))
                        file_count += 1

        conn.commit()
        conn.close()
        self.lbl_status.config(text=f"状態: 記憶完了 ({file_count}ファイル)", fg="green")
        messagebox.showinfo("完了", f"{file_count} 個のファイルを高速データベースに記憶しました。")

    def add_keyword_and_search(self, word):
        current = self.ent_keyword.get().strip()
        if word not in current:
            self.ent_keyword.delete(0, tk.END)
            self.ent_keyword.insert(0, f"{current} {word}".strip())
            self.run_search()

    def open_file(self, filepath):
        try:
            if os.path.exists(filepath):
                os.startfile(filepath)
            else:
                messagebox.showerror("エラー", "ファイルが見つからないか、移動・削除されています。")
        except Exception as e:
            messagebox.showerror("エラー", f"ファイルを開けませんでした: {e}")

    # 👇 指定場所への別名保存（ダウンロード）機能
    def save_file_to(self, filepath):
        if not os.path.exists(filepath):
            messagebox.showerror("エラー", "元ファイルが見つかりません。")
            return
        filename = os.path.basename(filepath)
        save_path = filedialog.asksaveasfilename(initialfile=filename, title="ファイルを別名で保存（コピー）")
        if save_path:
            try:
                shutil.copy2(filepath, save_path)
                messagebox.showinfo("成功", f"ファイルを保存しました:\n{save_path}")
            except Exception as e:
                messagebox.showerror("エラー", f"保存に失敗しました: {e}")

    def make_toggle_cmd(self, pf, btn):
        def toggle():
            if pf.winfo_ismapped():
                pf.pack_forget()
                btn.config(text="＋ 詳細")
            else:
                pf.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
                btn.config(text="－ 閉じる")
        return toggle

    def run_search(self):
        kw_input = self.ent_keyword.get().strip()
        if not kw_input:
            return
        if not os.path.exists(DB_FILE):
            messagebox.showwarning("未準備", "先に「① 資料フォルダを選択して記憶」ボタンを押してください。")
            return

        keywords = [w for w in kw_input.split() if len(w) > 0]
        conn = sqlite3.connect(DB_FILE)
        c = conn.cursor()

        sort_mode = self.sort_var.get()
        if "新着順" in sort_mode:
            order_sql = " ORDER BY mtime DESC"
        elif "古い順" in sort_mode:
            order_sql = " ORDER BY mtime ASC"
        else:
            order_sql = ""

        fts_words = [f'"{w}"' for w in keywords if len(w) >= 3]
        if fts_words:
            query = " AND ".join(fts_words)
            sql = f"SELECT filename, location, filepath, content, mtime FROM docs WHERE docs MATCH ?{order_sql} LIMIT 50"
            rows = c.execute(sql, (query,)).fetchall()
            for w in keywords:
                if len(w) < 3:
                    rows = [r for r in rows if w.lower() in r[3].lower()]
        else:
            like_clauses = " AND ".join(["content LIKE ?" for _ in keywords])
            params = [f"%{w}%" for w in keywords]
            sql = f"SELECT filename, location, filepath, content, mtime FROM docs WHERE {like_clauses}{order_sql} LIMIT 50"
            rows = c.execute(sql, params).fetchall()

        conn.close()

        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        self.txt_summary.delete("1.0", tk.END)

        if not rows:
            self.txt_summary.insert(tk.END, "該当する資料が見つかりませんでした。")
            return

        self.analyze_and_propose(keywords, rows)

        # ファイルごとにグループ化してカード表示
        file_dict = {}
        for fname, loc, fpath, content, mtime in rows:
            if fpath not in file_dict:
                file_dict[fpath] = {"filename": fname, "mtime": mtime, "snippets": []}
            file_dict[fpath]["snippets"].append((loc, content))

        for fpath, data in file_dict.items():
            date_str = datetime.fromtimestamp(data["mtime"]).strftime('%Y-%m-%d %H:%M') if data["mtime"] > 0 else "不明"
            
            card = tk.LabelFrame(self.scrollable_frame, bg="white", padx=5, pady=5)
            card.pack(fill=tk.X, padx=5, pady=4)

            header_frame = tk.Frame(card, bg="white")
            header_frame.pack(fill=tk.X)

            # 👇 折りたたみボタン
            toggle_btn = tk.Button(header_frame, text="＋ 詳細", width=6, bg="#e0e0e0", relief=tk.FLAT)
            toggle_btn.pack(side=tk.LEFT, padx=2)

            info_text = f"📄 {data['filename']}  (更新: {date_str})  パス: {fpath}"
            lbl_info = tk.Label(header_frame, text=info_text, bg="white", font=("Meiryo", 9, "bold"), anchor="w")
            lbl_info.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)

            btn_open = tk.Button(header_frame, text="📁 開く", bg="#e8f5e9", relief=tk.FLAT, command=lambda p=fpath: self.open_file(p))
            btn_open.pack(side=tk.RIGHT, padx=2)

            btn_save = tk.Button(header_frame, text="📥 保存", bg="#e3f2fd", relief=tk.FLAT, command=lambda p=fpath: self.save_file_to(p))
            btn_save.pack(side=tk.RIGHT, padx=2)

            # プレビュー展開エリア（最初は非表示）
            preview_frame = tk.Frame(card, bg="#fafafa", padx=5, pady=5)
            
            txt_preview = scrolledtext.Text(preview_frame, height=6, bg="#ffffff", font=("Meiryo", 9), wrap=tk.WORD)
            txt_preview.pack(fill=tk.BOTH, expand=True)
            txt_preview.tag_config("highlight", background="#ffff00", foreground="#d50000")

            for loc, content in data["snippets"]:
                txt_preview.insert(tk.END, f"【{loc}】\n{content.strip()}\n\n" + "-"*50 + "\n")

            for kw in keywords:
                start_pos = "1.0"
                while True:
                    start_pos = txt_preview.search(kw, start_pos, stopindex=tk.END, nocase=True)
                    if not start_pos:
                        break
                    end_pos = f"{start_pos}+{len(kw)}c"
                    txt_preview.tag_add("highlight", start_pos, end_pos)
                    start_pos = end_pos
            txt_preview.config(state=tk.DISABLED)

            toggle_btn.config(command=self.make_toggle_cmd(preview_frame, toggle_btn))

    def analyze_and_propose(self, keywords, rows):
        all_text = "\n".join([r[3] for r in rows])
        candidates = re.findall(r'[一-龥ァ-ヶーa-zA-Z0-9]{2,8}', all_text)
        filtered = [w for w in candidates if w not in STOP_WORDS and not any(w in kw or kw in w for kw in keywords) and not w.isdigit()]
        most_common = [w for w, _ in Counter(filtered).most_common(8)]

        for widget in self.frame_chips.winfo_children():
            widget.destroy()
        if most_common:
            tk.Label(self.frame_chips, text="💡 関連ワード（クリックで追加）: ", bg="#fffde7", font=("Meiryo", 9)).pack(side=tk.LEFT)
            for word in most_common:
                btn = tk.Button(self.frame_chips, text=f"＋ {word}", bg="#e3f2fd", relief=tk.FLAT, command=lambda w=word: self.add_keyword_and_search(w))
                btn.pack(side=tk.LEFT, padx=2)

        scored_sentences = []
        seen = set()
        for fname, loc, _, content, _ in rows:
            lines = re.split(r'[\n。]', content)
            for line in lines:
                line = line.strip()
                if len(line) < 12 or len(line) > 180 or line in seen:
                    continue
                seen.add(line)

                kw_score = sum(3 for kw in keywords if kw.lower() in line.lower())
                sig_score = sum(2 for sig in SIGNAL_WORDS if sig in line)
                total_score = kw_score + sig_score

                if kw_score > 0 and total_score >= 4:
                    scored_sentences.append((total_score, f"📌 [{fname} ({loc})] {line}"))

        scored_sentences.sort(key=lambda x: x[0], reverse=True)
        self.txt_summary.delete("1.0", tk.END)
        if scored_sentences:
            self.txt_summary.insert(tk.END, "【自動抽出された重要ポイント・回答候補】\n")
            for _, text in scored_sentences[:4]:
                self.txt_summary.insert(tk.END, text + "\n")
        else:
            self.txt_summary.insert(tk.END, "※ ヒットした資料の中から、結論・回答に関する定型文は検出されませんでした。下のファイル一覧をご確認ください。")

if __name__ == "__main__":
    root = tk.Tk()
    app = UltraLightSearchAppV2(root)
    root.mainloop()
